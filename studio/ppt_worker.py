"""Execute the fixed upstream SVG-to-DrawingML engine in an isolated child."""
from pathlib import Path
import json
import sys
from .paths import SOURCES

out=Path(sys.argv[1]).resolve()
sys.path.insert(0,str(SOURCES/'hugohe3--ppt-master/skills/ppt-master/scripts'))
from svg_to_pptx.pptx_package.builder import create_pptx_with_native_svg
from pptx import Presentation

manifest=json.loads((out/'manifest.json').read_text())
svgs=sorted((out/'slides').glob('*.svg'))
notes={p.stem:manifest['beats'][i]['narration'] for i,p in enumerate(svgs)}
ok=create_pptx_with_native_svg(svgs,out/'presentation.pptx',resource_root=out,
    verbose=False,use_native_shapes=True,pptx_structure='flat',notes=notes,workers=1,transition='fade',
    conversion_trace_path=out/'pptx-trace.json',expected_viewbox='0 0 1280 720',
    doc_metadata={'title':manifest['title'],'creator':'Family Creative Studio'},primary_language='zh-TW')
if not ok:
    raise RuntimeError('PPT Master rejected native conversion')
presentation=Presentation(out/'presentation.pptx')
from pptx.util import Inches
from PIL import Image
visuals=json.loads((out/'visuals.json').read_text()) if (out/'visuals.json').exists() else {}
for index,item in visuals.items():
    poster=out/item['poster']
    with Image.open(poster) as img:w,h=img.size
    scale=min(720/w,420/h);w*=scale;h*=scale
    presentation.slides[int(index)].shapes.add_picture(str(poster),int((510+(720-w)/2)/1280*presentation.slide_width),int((90+(420-h)/2)/720*presentation.slide_height),width=int(w/1280*presentation.slide_width),height=int(h/720*presentation.slide_height))
presentation.save(out/'presentation.pptx')
def walk(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape,'shapes'):
            yield from walk(shape.shapes)
all_shapes=[shape for slide in presentation.slides for shape in walk(slide.shapes)]
counts={'slides':len(presentation.slides),'editable_shapes':sum(1 for s in all_shapes if s.shape_type!=13),
        'text_frames':sum(bool(s.has_text_frame) for s in all_shapes),
        'pictures':sum(s.shape_type==13 for s in all_shapes),'native_paths':sum(bool(s._element.xpath('.//a:custGeom')) for s in all_shapes)}
if counts['slides']!=len(svgs) or counts['text_frames']<len(svgs)*2:
    raise RuntimeError('Editable slide content missing')
(out/'pptx-check.json').write_text(json.dumps(counts,indent=2))
