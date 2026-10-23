"""Native SVG fitting, inheritance and final editable PPT packaging regressions."""
import io
import json
import os
from pathlib import Path
import sys
from types import SimpleNamespace
import wave
import zipfile
from xml.etree import ElementTree as ET

import pytest

from studio.paths import SOURCES
from studio.presentation_svg import native_presentation_svg
from studio.svg_assets import normalized_svg

SVG = '{http://www.w3.org/2000/svg}'
UPSTREAM = SOURCES / 'hugohe3--ppt-master/skills/ppt-master/scripts'


@pytest.fixture(autouse=True)
def pinned_ppt_converter(monkeypatch):
    if not (UPSTREAM / 'svg_to_pptx').is_dir():
        pytest.skip('Pinned PPT Master checkout is required')
    monkeypatch.syspath_prepend(str(UPSTREAM))


ARTWORK = b'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="10 20 640 480"
  font-family="DejaVu Sans" font-size="20" fill="#114488">
  <rect x="10" y="20" width="640" height="480" fill="#FFFFFF"/>
  <circle cx="330" cy="260" r="100"/>
  <g style="font-size:24px;letter-spacing:2px;fill:#123456">
    <text x="110" y="220">Editable vector</text>
  </g>
  <path d="M 50 400 L 150 400 L 100 350 Z" fill="#CC6600"/>
</svg>'''


def test_offset_viewbox_bakes_inherited_text_metrics_and_retains_vectors():
    from svg_to_pptx.drawingml.utils import project_transform_errors, parse_transform_matrix
    root = native_presentation_svg(ARTWORK)
    assert root.get('viewBox') == '0 0 1280 720'
    assert project_transform_errors(root) == []
    text = root.find('.//' + SVG + 'text')
    assert text.attrib == {'x': '165', 'y': '330', 'font-size': '36', 'letter-spacing': '3'}
    wrapper = next(node for node in root.iter() if text in list(node))
    assert parse_transform_matrix(wrapper.get('transform')) == (1, 0, 0, 1, 145, -30)
    assert root.find(SVG + 'circle') is not None
    assert root.find(SVG + 'path') is not None
    assert root.find('.//' + SVG + 'image') is None


def test_nested_rotate_scale_and_relative_tspan_use_original_inheritance():
    from svg_to_pptx.drawingml.utils import project_transform_errors
    root = native_presentation_svg(b'''<svg xmlns="http://www.w3.org/2000/svg"
      viewBox="0 0 1280 720" font-size="20">
      <g transform="translate(100 50) rotate(30) scale(2)" style="font-size:1.5em;letter-spacing:0.1em">
        <text x="10" y="20" stroke="#000000">Parent<tspan x="10" dy="2em"
          style="font-size:2em;fill:#FF0000">Child</tspan></text>
      </g></svg>''')
    text = root.find('.//' + SVG + 'text')
    span = text.find(SVG + 'tspan')
    assert text.get('font-size') == '60' and text.get('letter-spacing') == '6'
    assert text.get('stroke-width') == '2'
    assert text.get('transform') == 'rotate(30)'
    assert span.get('font-size') == '120' and span.get('dy') == '240'
    assert span.get('letter-spacing') == '6'  # Inherited at its declaration site.
    assert 'font-size' not in span.get('style') and 'fill:#FF0000' in span.get('style')
    assert project_transform_errors(root) == []


@pytest.mark.parametrize('transform', ['scale(2 1)', 'matrix(1 0 .2 1 0 0)', 'scale(-1 1)'])
def test_unrepresentable_editable_text_never_silently_becomes_a_picture(transform):
    source = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">'
              f'<g transform="{transform}"><text x="10" y="20">Editable</text></g></svg>')
    with pytest.raises(ValueError, match='uniform scale and rotation'):
        native_presentation_svg(source.encode())


def test_baked_document_has_same_visible_pixels_as_original_fit():
    import cairosvg
    import numpy as np
    from PIL import Image
    # A rendering comparison verifies geometry/style meaning, not just our XML.
    before = cairosvg.svg2png(bytestring=normalized_svg(ARTWORK).encode())
    after = cairosvg.svg2png(bytestring=ET.tostring(native_presentation_svg(ARTWORK)))
    a = np.asarray(Image.open(io.BytesIO(before))).astype(float)
    b = np.asarray(Image.open(io.BytesIO(after))).astype(float)
    assert a.shape == b.shape == (720, 1280, 4)
    assert np.abs(a-b).mean() < .05


@pytest.mark.skipif(not os.environ.get('STUDIO_PPT_INTEGRATION'), reason='explicit native PPT worker integration')
def test_final_package_keeps_native_geometry_text_narration_and_valid_timing(tmp_path):
    from pptx import Presentation
    from studio.runner import Runner
    from studio.schema import Manifest
    for name in ('assets', 'slides', 'audio'):
        (tmp_path / name).mkdir()
    asset = 'a' * 64 + '.svg'
    (tmp_path / 'assets' / asset).write_bytes(ARTWORK)
    (tmp_path / 'slides/00.svg').write_text(normalized_svg(ARTWORK))
    with wave.open(str(tmp_path / 'audio/00.wav'), 'wb') as stream:
        stream.setnchannels(1); stream.setsampwidth(2); stream.setframerate(48000)
        stream.writeframes(b'\0\0' * 24000)
    manifest = Manifest.model_validate({'id': 'native-svg', 'title': 'Native vector', 'subtitle': 'Editable artwork',
        'slides': {'narration': True, 'autoplay': True, 'animation': 'fade'},
        'beats': [{'id': 'vector', 'title': 'Custom', 'body': 'Native', 'narration': 'Narration',
                   'duration': .5, 'svg_asset': asset}]})
    (tmp_path / 'manifest.json').write_text(manifest.model_dump_json())
    store = SimpleNamespace(root=tmp_path, job=lambda _: {'cancel': False})
    Runner(store, {'id': 'native-svg'}, tmp_path).run(
        [sys.executable, '-m', 'studio.ppt_worker', str(tmp_path)],
        label='ppt-native-svg', timeout=120, cwd=Path(__file__).resolve().parents[1])
    counts = json.loads((tmp_path / 'pptx-check.json').read_text())
    assert counts['slides'] == 1 and counts['native_paths'] == 1
    assert counts['narration_tracks'] == 1
    presentation = Presentation(tmp_path / 'presentation.pptx')

    def shapes(collection):
        for shape in collection:
            yield shape
            if hasattr(shape, 'shapes'):
                yield from shapes(shape.shapes)

    native = list(shapes(presentation.slides[0].shapes))
    text = next(shape for shape in native if shape.has_text_frame and shape.text == 'Editable vector')
    # PPT Master includes a small glyph-overhang allowance in text-box bounds.
    assert text.left == pytest.approx(310 * 9525, abs=2 * 9525)
    assert text.text_frame.paragraphs[0].runs[0].font.size.pt == 27
    assert str(text.text_frame.paragraphs[0].runs[0].font.color.rgb) == '123456'
    # The only raster is the upstream's tiny off-slide narration control icon.
    pictures = [shape for shape in native if shape.shape_type == 13]
    assert len(pictures) == 1 and pictures[0].left < 0 and pictures[0].top < 0
    ellipse = next(shape for shape in native if shape._element.xpath('.//a:prstGeom[@prst="ellipse"]'))
    assert (ellipse.left, ellipse.top, ellipse.width, ellipse.height) == tuple(
        value * 9525 for value in (490, 210, 300, 300))
    with zipfile.ZipFile(tmp_path / 'presentation.pptx') as package:
        slide = package.read('ppt/slides/slide1.xml').decode()
        assert 'prst="ellipse"' in slide and '<a:custGeom' in slide
        assert 'advTm="500"' in slide
        assert len([p for p in package.namelist() if p.startswith('ppt/media/') and p.endswith('.wav')]) == 1
        assert [p for p in package.namelist() if p.startswith('ppt/media/') and p.endswith('.png')] == ['ppt/media/narration_poster.png']


@pytest.mark.skipif(not os.environ.get('STUDIO_PPT_INTEGRATION'), reason='explicit native PPT worker integration')
def test_translated_imported_canvas_rectangle_remains_an_editable_shape(tmp_path):
    from pptx import Presentation
    from studio.runner import Runner
    from studio.schema import Manifest
    (tmp_path/'assets').mkdir();(tmp_path/'slides').mkdir()
    asset='b'*64+'.svg'
    artwork=b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><rect width="1280" height="720" fill="#FF0000" transform="translate(200 0)"/><text x="220" y="200">Preserve placement</text></svg>'
    (tmp_path/'assets'/asset).write_bytes(artwork)
    manifest=Manifest.model_validate({'id':'translated','title':'Translated vector','subtitle':'Native placement',
        'slides':{'narration':False,'animation':'fade'},
        'beats':[{'id':'vector','title':'Vector','body':'Native','narration':'','svg_asset':asset,'visual':{'audio':'silent'}}]})
    (tmp_path/'manifest.json').write_text(manifest.model_dump_json())
    store=SimpleNamespace(root=tmp_path,job=lambda _:{'cancel':False})
    Runner(store,{'id':'translated'},tmp_path).run([sys.executable,'-m','studio.ppt_worker',str(tmp_path)],
        label='ppt-translated',timeout=120,cwd=Path(__file__).resolve().parents[1])
    slide=Presentation(tmp_path/'presentation.pptx').slides[0]
    def shapes(collection):
        for shape in collection:
            yield shape
            if hasattr(shape,'shapes'):yield from shapes(shape.shapes)
    rect=next(s for s in shapes(slide.shapes) if s._element.xpath('.//a:prstGeom[@prst="rect"]'))
    assert (rect.left,rect.top,rect.width,rect.height)==tuple(x*9525 for x in (200,0,1280,720))
    assert str(rect.fill.fore_color.rgb)=='FF0000'
    assert not slide._element.xpath('./p:cSld/p:bg')
